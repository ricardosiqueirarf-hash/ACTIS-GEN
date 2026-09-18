"""Deterministic Excel (.xlsx) runtime tools for ACTIS GEN agents."""
from __future__ import annotations

import json
import os
import re
from pathlib import Path
from typing import Annotated, Any

from meuharness.artifacts import register_artifact

EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _libs():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.workbook.properties import CalcProperties
    except ImportError as exc:
        raise RuntimeError("Runtime Excel indisponível. Instale as dependências do ACTIS GEN.") from exc
    return {
        "Workbook": Workbook,
        "load_workbook": load_workbook,
        "BarChart": BarChart,
        "LineChart": LineChart,
        "PieChart": PieChart,
        "Reference": Reference,
        "Alignment": Alignment,
        "Border": Border,
        "Font": Font,
        "PatternFill": PatternFill,
        "Side": Side,
        "coordinate_to_tuple": coordinate_to_tuple,
        "range_boundaries": range_boundaries,
        "Table": Table,
        "TableStyleInfo": TableStyleInfo,
        "CalcProperties": CalcProperties,
    }


def _dump(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)


def _json_value(raw: str, label: str) -> Any:
    try:
        return json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{label} precisa ser JSON válido.") from exc


def _root() -> Path:
    configured = os.getenv("ACTIS_SPREADSHEET_ROOT", "").strip()
    return Path(configured).expanduser().resolve() if configured else Path.home().resolve()


def _safe_xlsx(path: str, *, must_exist: bool = False) -> Path:
    value = str(path or "").strip()
    if not value:
        raise ValueError("Caminho da planilha é obrigatório.")
    candidate = Path(value).expanduser()
    if not candidate.is_absolute():
        candidate = _root() / candidate
    candidate = candidate.resolve(strict=False)
    root = _root()
    if candidate != root and root not in candidate.parents:
        raise PermissionError(f"Planilhas só podem ser acessadas dentro de {root}.")
    if candidate.suffix.lower() != ".xlsx":
        raise ValueError("A skill Excel aceita apenas arquivos .xlsx.")
    if must_exist and not candidate.is_file():
        raise FileNotFoundError(str(candidate))
    return candidate


def _sheet(wb: Any, name: str | None):
    if name:
        if name not in wb.sheetnames:
            raise ValueError(f"Aba não encontrada: {name}")
        return wb[name]
    return wb.active


def _set_recalc(wb: Any, libs: dict[str, Any]) -> None:
    wb.calculation = libs["CalcProperties"](calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)


def _write_matrix(ws: Any, start: str, values: list[list[Any]], libs: dict[str, Any]) -> int:
    start_row, start_col = libs["coordinate_to_tuple"](start)
    written = 0
    for r_offset, row in enumerate(values):
        if not isinstance(row, list):
            raise TypeError("Cada linha precisa ser uma lista JSON.")
        for c_offset, value in enumerate(row):
            ws.cell(start_row + r_offset, start_col + c_offset, value=value)
            written += 1
    return written


def _apply_style(ws: Any, cell_range: str, style: dict[str, Any], libs: dict[str, Any]) -> None:
    min_col, min_row, max_col, max_row = libs["range_boundaries"](cell_range)
    border_cfg = style.get("border") or {}
    side = libs["Side"](
        style=str(border_cfg.get("style") or "thin") if border_cfg else None,
        color=str(border_cfg.get("color") or "D9D9D9").replace("#", "") if border_cfg else None,
    )
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            if any(key in style for key in ("bold", "italic", "font_size", "font_color")):
                cell.font = libs["Font"](
                    name=cell.font.name,
                    size=style.get("font_size", cell.font.sz),
                    bold=style.get("bold", cell.font.bold),
                    italic=style.get("italic", cell.font.italic),
                    color=str(style.get("font_color") or "").replace("#", "") or cell.font.color,
                )
            if style.get("fill"):
                cell.fill = libs["PatternFill"]("solid", fgColor=str(style["fill"]).replace("#", ""))
            if any(key in style for key in ("horizontal", "vertical", "wrap_text")):
                cell.alignment = libs["Alignment"](
                    horizontal=style.get("horizontal", cell.alignment.horizontal),
                    vertical=style.get("vertical", cell.alignment.vertical),
                    wrap_text=style.get("wrap_text", cell.alignment.wrap_text),
                )
            if style.get("number_format"):
                cell.number_format = str(style["number_format"])
            if border_cfg:
                cell.border = libs["Border"](left=side, right=side, top=side, bottom=side)


def _add_chart(ws: Any, operation: dict[str, Any], libs: dict[str, Any]) -> None:
    chart_type = str(operation.get("chart_type") or "bar").lower()
    chart_cls = {
        "bar": libs["BarChart"], "line": libs["LineChart"], "pie": libs["PieChart"],
    }.get(chart_type)
    if chart_cls is None:
        raise ValueError("chart_type deve ser bar, line ou pie.")
    data_range = str(operation.get("data_range") or "").strip()
    if not data_range:
        raise ValueError("add_chart exige data_range.")
    min_col, min_row, max_col, max_row = libs["range_boundaries"](data_range)
    chart = chart_cls()
    chart.title = str(operation.get("title") or "")
    chart.add_data(
        libs["Reference"](ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row),
        titles_from_data=bool(operation.get("titles_from_data", True)),
    )
    categories = str(operation.get("categories_range") or "").strip()
    if categories:
        c1, r1, c2, r2 = libs["range_boundaries"](categories)
        chart.set_categories(libs["Reference"](ws, min_col=c1, min_row=r1, max_col=c2, max_row=r2))
    chart.height = float(operation.get("height") or 7)
    chart.width = float(operation.get("width") or 12)
    ws.add_chart(chart, str(operation.get("anchor") or "H2"))


def _apply_operation(wb: Any, operation: dict[str, Any], libs: dict[str, Any]) -> dict[str, Any]:
    op = str(operation.get("op") or "").strip().lower()
    if op == "add_sheet":
        name = str(operation.get("name") or "Planilha").strip()[:31]
        if name in wb.sheetnames:
            raise ValueError(f"Aba já existe: {name}")
        wb.create_sheet(name)
        return {"op": op, "sheet": name}
    if op == "rename_sheet":
        ws = _sheet(wb, str(operation.get("sheet") or ""))
        new_name = str(operation.get("name") or "").strip()[:31]
        if not new_name:
            raise ValueError("rename_sheet exige name.")
        ws.title = new_name
        return {"op": op, "sheet": new_name}
    if op == "delete_sheet":
        if len(wb.sheetnames) <= 1:
            raise ValueError("Não é possível excluir a única aba.")
        ws = _sheet(wb, str(operation.get("sheet") or ""))
        name = ws.title
        wb.remove(ws)
        return {"op": op, "sheet": name}

    ws = _sheet(wb, str(operation.get("sheet") or "") or None)
    if op == "set":
        cell = str(operation.get("cell") or "").strip()
        if not cell:
            raise ValueError("set exige cell.")
        ws[cell] = operation.get("value")
    elif op == "write_range":
        values = operation.get("values")
        if not isinstance(values, list):
            raise TypeError("write_range exige values como matriz JSON.")
        count = _write_matrix(ws, str(operation.get("start") or "A1"), values, libs)
        return {"op": op, "sheet": ws.title, "cells_written": count}
    elif op == "append_rows":
        rows = operation.get("rows")
        if not isinstance(rows, list):
            raise TypeError("append_rows exige rows como matriz JSON.")
        for row in rows:
            if not isinstance(row, list):
                raise TypeError("Cada linha precisa ser uma lista JSON.")
            ws.append(row)
        return {"op": op, "sheet": ws.title, "rows_appended": len(rows)}
    elif op == "insert_rows":
        ws.insert_rows(int(operation.get("index") or 1), int(operation.get("amount") or 1))
    elif op == "delete_rows":
        ws.delete_rows(int(operation.get("index") or 1), int(operation.get("amount") or 1))
    elif op == "insert_cols":
        ws.insert_cols(int(operation.get("index") or 1), int(operation.get("amount") or 1))
    elif op == "delete_cols":
        ws.delete_cols(int(operation.get("index") or 1), int(operation.get("amount") or 1))
    elif op == "merge":
        ws.merge_cells(str(operation.get("range") or ""))
    elif op == "unmerge":
        ws.unmerge_cells(str(operation.get("range") or ""))
    elif op == "freeze_panes":
        ws.freeze_panes = operation.get("cell") or None
    elif op == "auto_filter":
        ws.auto_filter.ref = str(operation.get("range") or "") or None
    elif op == "column_width":
        ws.column_dimensions[str(operation.get("column") or "A").upper()].width = float(operation.get("width") or 12)
    elif op == "row_height":
        ws.row_dimensions[int(operation.get("row") or 1)].height = float(operation.get("height") or 15)
    elif op == "style_range":
        cell_range = str(operation.get("range") or "").strip()
        style = operation.get("style")
        if not cell_range or not isinstance(style, dict):
            raise ValueError("style_range exige range e style.")
        _apply_style(ws, cell_range, style, libs)
    elif op == "add_table":
        ref = str(operation.get("range") or "").strip()
        if not ref:
            raise ValueError("add_table exige range.")
        name = re.sub(r"[^A-Za-z0-9_]", "_", str(operation.get("name") or "Table1"))
        table = libs["Table"](displayName=name, ref=ref)
        table.tableStyleInfo = libs["TableStyleInfo"](
            name=str(operation.get("style") or "TableStyleMedium2"),
            showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)
    elif op == "add_chart":
        _add_chart(ws, operation, libs)
    else:
        raise ValueError(f"Operação Excel não suportada: {op}")
    return {"op": op, "sheet": ws.title}


def _summary(path: Path, libs: dict[str, Any], max_rows: int = 20, max_cols: int = 12) -> dict[str, Any]:
    wb = libs["load_workbook"](path, data_only=False, keep_links=False)
    sheets = []
    for ws in wb.worksheets:
        preview = [
            list(row)
            for row in ws.iter_rows(
                min_row=1, max_row=min(ws.max_row, max(1, max_rows)),
                min_col=1, max_col=min(ws.max_column, max(1, max_cols)), values_only=True,
            )
        ]
        sheets.append({
            "name": ws.title, "rows": ws.max_row, "columns": ws.max_column,
            "tables": list(ws.tables.keys()), "charts": len(ws._charts),
            "merged_ranges": [str(item) for item in ws.merged_cells.ranges], "preview": preview,
        })
    return {"path": str(path), "sheets": sheets}


def build_excel_tools(
    scopes: set[str] | None = None, *, run_id: str | None = None, agent_id: str | None = None
) -> list[object]:
    """Build native Excel tools constrained by spreadsheet.read / spreadsheet.write scopes."""
    allowed = set(scopes or {"spreadsheet.read", "spreadsheet.write"})
    libs = _libs()
    tools: list[object] = []

    if "spreadsheet.write" in allowed:
        def excel_create(
            output_path: Annotated[str, "Novo .xlsx dentro do diretório do usuário"],
            workbook_json: Annotated[str, "JSON {sheets:[{name,rows,freeze_panes,auto_filter,column_widths,styles}]}"],
        ) -> str:
            """Crie uma planilha Excel real com dados, fórmulas e formatação."""
            target = _safe_xlsx(output_path)
            spec = _json_value(workbook_json, "workbook_json")
            if not isinstance(spec, dict):
                raise TypeError("workbook_json precisa ser um objeto JSON.")
            target.parent.mkdir(parents=True, exist_ok=True)
            wb = libs["Workbook"]()
            wb.remove(wb.active)
            sheets = spec.get("sheets") or [{"name": "Planilha1", "rows": []}]
            if not isinstance(sheets, list) or not sheets:
                raise ValueError("sheets precisa ter ao menos uma aba.")
            for sheet_spec in sheets:
                if not isinstance(sheet_spec, dict):
                    raise TypeError("Cada sheet precisa ser um objeto JSON.")
                ws = wb.create_sheet(str(sheet_spec.get("name") or "Planilha")[:31])
                rows = sheet_spec.get("rows") or []
                if not isinstance(rows, list):
                    raise TypeError("rows precisa ser uma matriz JSON.")
                for row in rows:
                    if not isinstance(row, list):
                        raise TypeError("Cada linha precisa ser uma lista JSON.")
                    ws.append(row)
                if sheet_spec.get("freeze_panes"):
                    ws.freeze_panes = str(sheet_spec["freeze_panes"])
                if sheet_spec.get("auto_filter"):
                    ws.auto_filter.ref = str(sheet_spec["auto_filter"])
                for col, width in dict(sheet_spec.get("column_widths") or {}).items():
                    ws.column_dimensions[str(col).upper()].width = float(width)
                for style_item in sheet_spec.get("styles") or []:
                    if isinstance(style_item, dict) and style_item.get("range") and isinstance(style_item.get("style"), dict):
                        _apply_style(ws, str(style_item["range"]), style_item["style"], libs)
            _set_recalc(wb, libs)
            wb.save(target)
            artifact = None
            if run_id:
                artifact = register_artifact(target, run_id=run_id, agent_id=agent_id, mime_type=EXCEL_MIME)
            return _dump({"ok": True, "artifact_id": artifact.get("id") if artifact else None, **_summary(target, libs)})

        def excel_edit(
            path: Annotated[str, "Caminho do .xlsx existente"],
            operations_json: Annotated[str, "Lista JSON de operações set/write_range/append_rows/sheets/rows/cols/style/table/chart"],
        ) -> str:
            """Edite um workbook existente preservando o restante da estrutura."""
            target = _safe_xlsx(path, must_exist=True)
            operations = _json_value(operations_json, "operations_json")
            if not isinstance(operations, list) or not operations:
                raise ValueError("operations_json precisa ser uma lista não vazia.")
            wb = libs["load_workbook"](target, data_only=False, keep_links=False)
            results = []
            for index, operation in enumerate(operations):
                if not isinstance(operation, dict):
                    raise TypeError(f"Operação {index} precisa ser um objeto JSON.")
                try:
                    results.append(_apply_operation(wb, operation, libs))
                except Exception as exc:
                    raise ValueError(f"Falha na operação {index} ({operation.get('op')}): {exc}") from exc
            _set_recalc(wb, libs)
            wb.save(target)
            artifact = None
            if run_id:
                artifact = register_artifact(target, run_id=run_id, agent_id=agent_id, mime_type=EXCEL_MIME)
            return _dump({"ok": True, "path": str(target), "artifact_id": artifact.get("id") if artifact else None, "operations": results})

        tools.extend([excel_create, excel_edit])

    if "spreadsheet.read" in allowed:
        def excel_inspect(
            path: Annotated[str, "Caminho do .xlsx"],
            sheet: Annotated[str, "Nome da aba; vazio = todas"] = "",
            cell_range: Annotated[str, "Range A1:F20; vazio = preview"] = "",
            max_rows: Annotated[int, "Máximo de linhas"] = 40,
            max_cols: Annotated[int, "Máximo de colunas"] = 20,
        ) -> str:
            """Leia estrutura, valores e fórmulas de uma planilha sem alterá-la."""
            target = _safe_xlsx(path, must_exist=True)
            wb = libs["load_workbook"](target, data_only=False, keep_links=False)
            if cell_range:
                ws = _sheet(wb, sheet or None)
                min_col, min_row, max_col, max_row = libs["range_boundaries"](cell_range)
                values = [[cell.value for cell in row] for row in ws.iter_rows(
                    min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
                )]
                return _dump({"path": str(target), "sheet": ws.title, "range": cell_range, "values": values})
            if sheet:
                ws = _sheet(wb, sheet)
                preview = [list(row) for row in ws.iter_rows(
                    min_row=1, max_row=min(ws.max_row, max(1, int(max_rows))),
                    min_col=1, max_col=min(ws.max_column, max(1, int(max_cols))), values_only=True,
                )]
                return _dump({"path": str(target), "sheet": ws.title, "rows": ws.max_row, "columns": ws.max_column,
                              "tables": list(ws.tables.keys()), "charts": len(ws._charts), "preview": preview})
            return _dump(_summary(target, libs, int(max_rows), int(max_cols)))

        def excel_validate(path: Annotated[str, "Caminho do .xlsx a validar"]) -> str:
            """Valide se o workbook abre e resuma fórmulas/estrutura antes da entrega."""
            target = _safe_xlsx(path, must_exist=True)
            wb = libs["load_workbook"](target, data_only=False, keep_links=False)
            issues: list[str] = []
            formulas = 0
            for ws in wb.worksheets:
                if ws.sheet_state != "visible":
                    issues.append(f"Aba oculta: {ws.title}")
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            formulas += 1
                for table_name in ws.tables.keys():
                    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", table_name):
                        issues.append(f"Nome de tabela potencialmente incompatível: {table_name}")
            return _dump({
                "ok": not issues, "path": str(target), "sheets": wb.sheetnames,
                "formula_count": formulas, "issues": issues,
                "note": "Fórmulas são preservadas e o workbook é marcado para recalcular ao abrir no Excel/LibreOffice.",
            })

        tools.extend([excel_inspect, excel_validate])

    # Backward-compatible names used by the inheritable skill_excel contract.
    def excel_create_workbook(output_path: str, workbook_json: str) -> str:
        if "spreadsheet.write" not in allowed:
            raise PermissionError("Scope spreadsheet.write não autorizado. Solicite ACTIS_APPROVAL_REQUEST: spreadsheet.write.")
        return excel_create(output_path, workbook_json)

    def excel_edit_workbook(path: str, operations_json: str) -> str:
        if "spreadsheet.write" not in allowed:
            raise PermissionError("Scope spreadsheet.write não autorizado. Solicite ACTIS_APPROVAL_REQUEST: spreadsheet.write.")
        return excel_edit(path, operations_json)

    def excel_inspect_workbook(path: str, sheet: str = "", cell_range: str = "", max_rows: int = 40, max_cols: int = 20) -> str:
        if "spreadsheet.read" not in allowed:
            raise PermissionError("Scope spreadsheet.read não autorizado. Solicite ACTIS_APPROVAL_REQUEST: spreadsheet.read.")
        return excel_inspect(path, sheet, cell_range, max_rows, max_cols)

    def excel_validate_workbook(path: str) -> str:
        if "spreadsheet.read" not in allowed:
            raise PermissionError("Scope spreadsheet.read não autorizado. Solicite ACTIS_APPROVAL_REQUEST: spreadsheet.read.")
        return excel_validate(path)

    tools.extend([excel_create_workbook, excel_edit_workbook, excel_inspect_workbook, excel_validate_workbook])
    return tools
